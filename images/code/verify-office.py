#!/usr/bin/env python3
"""构建期验证办公文件工具链（images/code/Dockerfile 5c 段调用，失败即构建失败）。

背景：B 端客户断网，call-tool 失败日志里最多的是 .doc/.xls/.wps 缺 olefile/xlrd/antiword、
`file` 命令缺失、markitdown 缺 docx extra。这里真实生成 docx/xlsx/pdf 各转换一遍，并逐个确认
系统命令存在，防「声称预装、实际缺依赖」。
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

# 1. Python 模块可 import
import fitz  # noqa: E402  pymupdf
import mammoth  # noqa: E402  markitdown[docx]
import msoffcrypto  # noqa: E402
import olefile  # noqa: E402
import xlrd  # noqa: E402
import docx  # noqa: E402
import openpyxl  # noqa: E402
from markitdown import MarkItDown  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

print("python libs:", "pymupdf", fitz.__doc__ or "", "| olefile", olefile.__version__, "| xlrd", xlrd.__version__,
      "| msoffcrypto", getattr(msoffcrypto, "__version__", "?"), "| mammoth", getattr(mammoth, "__version__", "?"))

# 2. 系统命令存在
missing = [c for c in ("file", "xxd", "antiword", "catdoc", "xls2csv", "catppt",
                       "pdftotext", "pdfinfo", "pdftoppm", "pdfimages") if shutil.which(c) is None]
if missing:
    sys.exit(f"缺系统命令: {missing}")

tmp = Path(tempfile.mkdtemp())
md = MarkItDown()

# 3. docx：python-docx 写 → markitdown 读
d = docx.Document(); d.add_paragraph("沙盒docx验证"); d.save(tmp / "t.docx")
assert "沙盒docx验证" in md.convert(str(tmp / "t.docx")).text_content, "markitdown 读 docx 失败"

# 4. xlsx：openpyxl 写 → markitdown 读
w = openpyxl.Workbook(); w.active["A1"] = "沙盒xlsx验证"; w.save(tmp / "t.xlsx")
assert "沙盒xlsx验证" in md.convert(str(tmp / "t.xlsx")).text_content, "markitdown 读 xlsx 失败"

# 5. pdf：reportlab 写 → pymupdf / markitdown / pdftotext 读
c = canvas.Canvas(str(tmp / "t.pdf")); c.drawString(100, 700, "sandbox pdf ok"); c.save()
assert "sandbox pdf ok" in fitz.open(str(tmp / "t.pdf"))[0].get_text(), "pymupdf 读 pdf 失败"
assert "sandbox pdf ok" in md.convert(str(tmp / "t.pdf")).text_content, "markitdown 读 pdf 失败"
out = subprocess.run(["pdftotext", str(tmp / "t.pdf"), "-"], capture_output=True, text=True, check=True).stdout
assert "sandbox pdf ok" in out, "pdftotext 失败"

# 6. file 命令能识别 zip 容器类 Office
out = subprocess.run(["file", str(tmp / "t.docx")], capture_output=True, text=True, check=True).stdout
assert "Word" in out or "Zip" in out or "OOXML" in out, f"file 识别 docx 异常: {out}"

shutil.rmtree(tmp, ignore_errors=True)
print("office toolchain ok: docx/xlsx/pdf 转换 + file/antiword/catdoc/poppler 命令齐全")
